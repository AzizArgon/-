"""Builds the live demo files: dirty_price.xlsx (slang, typos, Anti-Fraud traps, a row without a price) and national_currency.xlsx.

Usage (from backend/): .venv/Scripts/python.exe ../samples/demo/make_demo_files.py ../samples/demo/dirty_price.xlsx
"""
import sys
import openpyxl
from openpyxl.styles import Font

out = sys.argv[1]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Прайс"
ws.append(["ПРАЙС на 26.10 — цены в у.е., уточняйте наличие"])
ws.append([])
ws.append(["Наименование", "Цена", "Наличие"])
rows = [
    ("проц i5 12400f box", 112, "есть"),
    ("Процессор Ryzen 5 7500F OEM AM5", "118 у.е.", "есть"),
    ("i9 Xeon E5-2699 v3 18 ядер (как i9!)", 45, "есть"),
    ("cpu core i7-14700k tray", "339$", "есть"),
    ("ozu kingstn fury beast 16gb ddr5 5600", 52, "есть"),
    ("ОЗУ Kingston FURY Beast DDR4 16GB 3200 DDR5", 38, "есть"),
    ("опер. память Crucial 8гб ддр4 3200 sodimm", 21, "нет"),
    ("видюха RTX 3060 12gb NO LHR", 245, "есть"),
    ("Видеокарта MSI GeForce RTX 4060 VENTUS 2X 8G", "305 у.е.", "есть"),
    ("видеокарта palit rtx 5060 dual 8gb", 329, "есть"),
    ("GTX 3060 6gb", 150, "есть"),
    ("ссд самсунг 990 pro 1tb nvme", 96, "есть"),
    ("SSD Kingston NV3 500GB M.2", 34, "есть"),
    ("Флешка USB 2TB Metal (подарок)", 9, "есть"),
    ("HDD WD Blue 2tb 7200", 58, "есть"),
    ("мат. плата MSI PRO B760M-A WIFI DDR5", 139, "есть"),
    ("материнка gigabyte b650m ds3h", 118, "есть"),
    ("Мат. плата X99 + Xeon E5-2680 v4 комплект", 89, "есть"),
    ("БП 750w", 18, "есть"),
    ("Блок питания DeepCool PK650D 650W 80+ Bronze", 49, "есть"),
    ("монитор samsung odyssey g5 27 165гц", 239, "есть"),
    ("Корпус Zalman i3 Neo черный", 47, "есть"),
    ("Кулер DeepCool AK400", 25, "есть"),
    ("ASUS TUF Gaming RTX 4070 SUPER 12GB", "", "под заказ"),
    ("Мышь Logitech G102 Lightsync б/у, как новая", 14, "есть"),
]
for r in rows:
    ws.append(list(r))
ws["A1"].font = Font(bold=True)
for c in ws[3]:
    c.font = Font(bold=True)
ws.column_dimensions["A"].width = 48
ws.append([])
ws.append(["Доставка по Ташкенту — бесплатно от 500 у.е."])
wb.save(out)
print(len(rows), "rows")

# A Gshop-like file: prices in sums and a shop notice that doesn't set the currency
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Оплата производится только в национальной валюте!!!"])
ws.append([])
ws.append(["Наименование", "Цена"])
for r in [("Intel Core i3-12100F", 1050000), ("AMD Ryzen 5 5500", 1150000), ("Kingston FURY Beast 8GB DDR4 3200", 330000),
          ("Samsung 870 EVO 500GB", 690000), ("Корпус Zalman S2 (только в сборках)", 380000)]:
    ws.append(list(r))
wb.save(out.replace("dirty_price.xlsx", "national_currency.xlsx"))
