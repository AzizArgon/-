"""Sheets with several "name | price" tables side by side (Gshop): split into vertical blocks.

Each block has its own section headers ("ПРОЦЕССОРЫ INTEL", sometimes two levels: "ПЕРИФЕРИЯ JBL" -> "НАУШНИКИ"),
so sections are followed in every block apart. Shop notices ("Оплата ... в национальной валюте!!!") are not goods.
"""

import io
import re
from dataclasses import dataclass

import openpyxl

from app.services.price_cleaner import parse_price

_NOTICE = re.compile(r"оплат|валют|курс|цены указаны|!!!", re.IGNORECASE)
_HEADER_WORDS = re.compile(r"^(наименование|название|товар|цена|price|name)$", re.IGNORECASE)
# A header row of a table: "Наименование", "Наименование товара", "Цена, $"
_HEADER_CELL = re.compile(r"^(наименовани|название|товар|цена|price|name)", re.IGNORECASE)
NOTE_TEXT = {"only_in_builds": "только в сборках"}
_ONLY_IN_BUILDS = re.compile(r"\s*\(\s*только в сборках\s*\)", re.IGNORECASE)
MIN_PRICE_SHARE = 0.3  # a price column: at least 30% of its filled cells are numbers


@dataclass
class BlockItem:
    block: int
    row: int
    section: str | None
    name: str
    price: float | None
    note: str | None = None  # "only_in_builds": the price holds only inside a PC build
    price_text: str = ""  # the price cell as written ("1.250.000 сум"): the currency stays with it


def _text(value) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip() if value is not None else ""


def _number(value) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return parse_price(_text(value))  # "58", "1 250 000", "1.250.000 сум", "110$"


def find_blocks(rows: list[list]) -> list[tuple[int, int]]:
    """(name column, price column) pairs. Price columns hold mostly numbers; the name is the column to the left."""
    width = max((len(r) for r in rows), default=0)
    pairs = []
    for col in range(1, width):
        filled = [r[col] for r in rows if col < len(r) and _text(r[col])]
        names = [r[col - 1] for r in rows if _text(r[col - 1]) and _number(r[col - 1]) is None]
        numbers = sum(1 for v in filled if _number(v) is not None)
        if filled and numbers / len(filled) >= MIN_PRICE_SHARE and len(names) >= numbers:
            if not pairs or pairs[-1][1] != col - 1:
                pairs.append((col - 1, col))
    return pairs


def _is_section(name: str) -> bool:
    # Headers are written in capitals and have no price: "ПРОЦЕССОРЫ INTEL", "LG", "STREAM DECK", "TOSHIBA 3TB".
    # A capitalised item without a price has a model code among several words: "MSI PRO B660M-E LGA 1700 DDR4"
    words = name.split()
    has_code = any(re.search(r"\d", w) and re.search(r"[A-ZА-ЯЁ]", w) for w in words)
    return (len(name) <= 60 and not re.search(r"[a-zа-яё]", name) and re.search(r"[A-ZА-ЯЁ]", name) is not None
            and not (has_code and len(words) >= 3))


def block_items(rows: list[list]) -> tuple[list[BlockItem], list[str]]:
    """Goods of all blocks (top to bottom, block by block) and shop notices found on the sheet."""
    items, notices = [], []
    for b, (name_col, price_col) in enumerate(find_blocks(rows)):
        section, prev_header = None, None
        for i, row in enumerate(rows):
            name = _text(row[name_col]) if name_col < len(row) else ""
            price = _number(row[price_col]) if price_col < len(row) else None
            if not name or _HEADER_WORDS.match(name):
                prev_header = None
                continue
            if price is None and _NOTICE.search(name):
                notices.append(name)
                continue
            if price is None and _is_section(name):
                # Two headers in a row: brand -> type ("ПЕРИФЕРИЯ JBL" / "НАУШНИКИ")
                section = f"{prev_header} / {name}" if prev_header else name
                prev_header = name if not prev_header else prev_header
                continue
            prev_header = None
            note = "only_in_builds" if _ONLY_IN_BUILDS.search(name) else None
            items.append(BlockItem(b, i, section, _ONLY_IN_BUILDS.sub("", name).strip(), price, note,
                                   _text(row[price_col]) if price_col < len(row) else ""))
    return items, list(dict.fromkeys(notices))


def _fill_merged(ws) -> None:
    """A cell merged down over several rows (a brand next to its models) belongs to every one of those rows;
    a cell merged only across (a section title over the table) stays a single cell."""
    for rng in list(ws.merged_cells.ranges):
        value = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        if rng.max_row > rng.min_row:
            for r in range(rng.min_row, rng.max_row + 1):
                ws.cell(r, rng.min_col).value = value


def _is_header(row: list) -> bool:
    return sum(bool(_HEADER_CELL.match(_text(c))) for c in row) >= 1 and not any(_number(c) for c in row)


def _regions(rows: list[list]) -> list[tuple[list | None, list[list]]]:
    """(header row, rows below it up to the next header): every region may have its own table layout."""
    regions: list[tuple[list | None, list[list]]] = [(None, [])]
    for row in rows:
        if _is_header(row):
            regions.append((row, []))
        else:
            regions[-1][1].append(row)
    return [r for r in regions if r[0] is not None or r[1]]


def prepare_workbook(data: bytes) -> tuple[bytes, list[str]]:
    """Excel price list before reading: merged cells filled, shop notices taken out (returned as hints for the
    seller), and a region with several "name | price" tables side by side turned into one table
    "Раздел | Наименование | Цена | Примечание" (sections followed in every block apart)."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    notices: list[str] = []
    for ws in wb.worksheets:
        _fill_merged(ws)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "|" in cell.value:
                    cell.value = cell.value.replace("|", "/")  # "|" inside a cell would split the table row
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        out: list[list] = []
        for header, part in _regions(rows):
            kept = []
            for row in part:
                filled = [c for c in row if _text(c)]
                if len(filled) == 1 and _NOTICE.search(_text(filled[0])) and _number(filled[0]) is None:
                    notices.append(_text(filled[0]))  # "Оплата ... в национальной валюте!!!" is not a section
                else:
                    kept.append(row)
            if len(find_blocks(kept)) >= 2:
                items, found = block_items(kept)
                notices += found
                out.append(["Раздел", "Наименование", "Цена", "Примечание"])
                out += [[it.section, it.name, it.price_text or None, NOTE_TEXT.get(it.note)] for it in items]
            else:
                out += ([header] if header else []) + kept
        if out != rows:
            ws.delete_rows(1, ws.max_row)
            for row in out:
                ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), list(dict.fromkeys(notices))
