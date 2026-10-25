"""Phase 7: dirty price files (merged cells, several tables, prices with currencies, PDF) and card matching."""

import io

import openpyxl
import pytest
from fpdf import FPDF

from app.services.catalog import make_match_key
from app.services.importer import file_to_markdown
from app.services.price_cleaner import clean_markdown, parse_price, price_currency
from app.services.spec_normalizer import canonical_brand, normalize_specs


def _xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(data: bytes, name: str, notices: list | None = None):
    return clean_markdown(file_to_markdown(data, name, notices))


@pytest.mark.parametrize("cell, price, currency", [
    ("1 250 000 сум", 1_250_000, "UZS"), ("1.250.000 сум", 1_250_000, "UZS"), ("140.000 сум", 140_000, "UZS"),
    ("110$", 110, "USD"), ("110 у.е.", 110, "USD"), ("110 y.e", 110, "USD"), ("110,00", 110, None), ("61.5", 61.5, None),
])
def test_price_cell_formats_and_currency(cell, price, currency):
    assert (parse_price(cell), price_currency(cell)) == (price, currency)


def test_merged_brand_cell_belongs_to_every_row_and_title_stays_a_section():
    def build(ws):
        ws.append(["Бренд", "Модель", "Цена"])
        ws.append(["ОПЕРАТИВНАЯ ПАМЯТЬ"])
        ws.merge_cells("A2:C2")
        ws.append(["Kingston", "DDR4 16GB 3200MHz", 40])
        ws.append([None, "DDR5 32GB 6000MHz", 120])
        ws.merge_cells("A3:A4")
    rows = _rows(_xlsx(build), "p.xlsx")
    assert [(r.section, r.name, r.price) for r in rows] == [
        ("ОПЕРАТИВНАЯ ПАМЯТЬ", "Kingston DDR4 16GB 3200MHz", 40), ("ОПЕРАТИВНАЯ ПАМЯТЬ", "Kingston DDR5 32GB 6000MHz", 120)]


def test_side_by_side_tables_notice_and_build_only_note():
    def build(ws):
        ws.append(["Оплата только в национальной валюте!!!"])
        ws.append(["Наименование", "Цена", None, "Наименование", "Цена"])
        ws.append(["ПРОЦЕССОРЫ", None, None, "КОРПУСА", None])
        ws.append(["Intel Core i5-12400F", "1 650 000", None, "Lian Li LANCOOL 205", "1.420.000 сум"])
        ws.append(["AMD Ryzen 5 5600", 1_500_000, None, "Corsair SF600 (Только в сборках)", 1_100_000])
        ws.append([])
        ws.append(["№", "Наименование", "Артикул", "Цена"])
        ws.append([1, "Kingston Fury DDR4 16GB", "KF432C16BB/16", "480 000"])
    notices: list[str] = []
    rows = _rows(_xlsx(build), "p.xlsx", notices)
    assert notices == ["Оплата только в национальной валюте!!!"]
    got = [(r.section, r.name, r.price, r.currency, r.note) for r in rows]
    assert got == [
        ("ПРОЦЕССОРЫ", "Intel Core i5-12400F", 1_650_000, None, None),
        ("ПРОЦЕССОРЫ", "AMD Ryzen 5 5600", 1_500_000, None, None),
        ("КОРПУСА", "Lian Li LANCOOL 205", 1_420_000, "UZS", None),
        ("КОРПУСА", "Corsair SF600", 1_100_000, None, "только в сборках"),
        ("", "Kingston Fury DDR4 16GB", 480_000, None, None),
    ]


def test_pdf_table_is_read_row_by_row():
    pdf = FPDF()
    pdf.add_font("arial", fname=r"C:\Windows\Fonts\arial.ttf")
    pdf.add_font("arial", style="B", fname=r"C:\Windows\Fonts\arialbd.ttf")
    pdf.set_font("arial", size=9)
    pdf.add_page()
    with pdf.table(col_widths=(150, 40)) as table:
        for row in (["Наименование", "Цена"], ["Intel Core i5-12400F OEM", "114$"], ["SSD Kingston 1TB NVMe NV3", "1 680 000 сум"]):
            table.row(row)
    rows = _rows(bytes(pdf.output()), "p.pdf")
    assert [(r.name, r.price, r.currency) for r in rows] == [
        ("Intel Core i5-12400F OEM", 114, "USD"), ("SSD Kingston 1TB NVMe NV3", 1_680_000, "UZS")]


def test_text_line_with_the_price_at_the_end():
    rows = clean_markdown("Прайс-лист\nIntel Core i5-12400F OEM 114$\nSSD Kingston 1TB NV3 — 1 680 000 сум")
    assert [(r.name, r.price, r.currency) for r in rows] == [
        ("Intel Core i5-12400F OEM", 114, "USD"), ("SSD Kingston 1TB NV3", 1_680_000, "UZS")]


def test_match_key_ignores_descriptions_lookalikes_and_glued_chips():
    rule = make_match_key("gpu", "Zotac GeForce RTX 5060 Ti TWIN EDGE 8GB", "Zotac")
    assert make_match_key("gpu", "Видеокарта Zotac RTX5060Ti TWIN EDGE 8GB GDDR7 128bit", "Zotac") == rule
    monitor = make_match_key("monitor", "ASUS TUF Gaming VG249QM1A 24\"", "ASUS")
    assert make_match_key("monitor", "Монитор ASUS TUF VG249QM1A 24\" Fast IPS 270Hz FHD 1920x1080", "ASUS") == monitor
    assert make_match_key("monitor", "Pixel PХG27Q 27\"", "Pixel") == make_match_key("monitor", "Pixel PXG27Q 27\"", "Pixel")
    cpu = make_match_key("cpu", "AMD Ryzen 7 9800X3D OEM", "AMD")
    assert make_match_key("cpu", "AMD Ryzen 7 9800X3D 4.7 GHz Granite Ridge", "AMD", {"pack": "OEM"}) == cpu
    assert make_match_key("ram", "Kingston Fury DDR4 16 ГБ 3200", "Kingston") == make_match_key("ram", "Kingston Fury DDR4 16GB 3200MHz", "Kingston")


def test_match_key_keeps_what_tells_products_apart():
    board = lambda ddr: make_match_key("motherboard", "Gigabyte B860M Gaming WiFi6", "Gigabyte", {"memory_type": ddr})  # noqa: E731
    assert board("DDR5") != board("DDR4")
    assert make_match_key("gpu", "RTX 4060 Ti", "MSI") != make_match_key("gpu", "RTX 4060", "MSI")
    assert make_match_key("ram", "Kingston DDR4 16GB 3200MHz") != make_match_key("ram", "Kingston DDR4 16GB 3600MHz")


@pytest.mark.parametrize("written, brand", [
    ("SAPPHRIE", "Sapphire"), ("ASSUS", "ASUS"), ("Lexsar", "Lexar"), ("Gigabyet", "Gigabyte"),
    ("Apacer", "Apacer"), ("Thermalright", "Thermalright"), ("MCI", "MCI"), ("Peladin", "Peladin"), ("Vissonic", "Vissonic"),
    ("Samsnug", "Samsung"), ("Zotak", "Zotac"), ("Hkisemi", "Hiksemi"),
])
def test_brand_typos(written, brand):
    assert canonical_brand(written) == brand


def test_refresh_rate_is_not_cut_and_cooler_is_not_a_paste():
    assert normalize_specs("monitor", "LG 24GN60R IPS 1440Hz").get("_suspect") == {"refresh_rate": 1440}
    assert normalize_specs("cooling", "Кулер AG300 от DeepCool | Заводская термопаста")["type"] == "air"
    assert normalize_specs("cooling", "Термопаста Arctic MX-4 для кулера")["type"] == "paste"


def test_chip_vendor_is_not_the_brand_of_a_card():
    from app.services.importer import _preview_row
    from app.services.price_cleaner import CleanRow
    row = lambda text: CleanRow(row=1, sheet="", section="", cells=[text], name=text, ai_text=text)  # noqa: E731
    ai = {"category": "gpu", "brand": "NVIDIA", "name": "NVIDIA GeForce RTX 3060 12GB"}
    assert _preview_row(row("rtx 3060 12gb"), ai, "ai")["brand"] is None
    assert _preview_row(row("NVIDIA RTX A2000 6GB"), {**ai, "name": "NVIDIA RTX A2000"}, "ai")["brand"] == "NVIDIA"
    typo = _preview_row(row("MB SAPPHRIE B850A"), {"category": "motherboard", "brand": "SAPPHRIE", "name": "SAPPHRIE PURE B850A"}, "ai")
    assert (typo["brand"], typo["name"]) == ("Sapphire", "Sapphire PURE B850A")


def test_ai_answer_in_another_shape_is_accepted():
    from app.services.ai_parser import GroqAnswer, _to_dicts
    answer = '{"items": [{"i": 1, "c": "ram", "n": "Kingston DDR4 8GB", "s": "unknown", "p": {"type": "DDR4", "capacity": "8GB"}, "f": ["weird"]}]}'
    item, = _to_dicts(GroqAnswer.model_validate_json(answer).items)
    assert (item["specs"], item["condition"], item["risks"]) == ({"type": "DDR4", "capacity": "8GB"}, "new", [])


def test_liquid_cooler_named_a_cooler_is_still_liquid():
    assert normalize_specs("cooling", "Кулер для процессора MSI MAG CORELIQUID 240R V2")["type"] == "liquid"


def test_row_without_identity_value_goes_to_the_only_card_that_has_it():
    from app.services.catalog import completion_index, unique_completion
    cards = [make_match_key("monitor", "Titan Army KG277S 27\"", "Titan Army", {"refresh_rate": 240}),
             make_match_key("monitor", "Redmi G27Q 27\"", "Redmi", {"refresh_rate": 240}),
             make_match_key("monitor", "Redmi G27Q 27\"", "Redmi", {"refresh_rate": 320})]
    index = completion_index(cards)
    assert unique_completion(make_match_key("monitor", "моник titan army kg277s 27 дюйм", None), index) == cards[0]
    assert unique_completion(make_match_key("monitor", "Redmi G27Q 27\"", "Redmi"), index) is None  # two models
    assert unique_completion(make_match_key("monitor", "Titan Army KG277S 27\"", "Titan Army", {"refresh_rate": 180}), index) is None


def test_pipe_inside_a_cell_and_no_mining():
    from app.services.antifraud import check_offer
    rows = _rows(_xlsx(lambda ws: [ws.append(r) for r in (["Наименование", "Цена"], ["EDIFIER R1080 BT AUX | Bluetooth 5.0", 75])]), "p.xlsx")
    assert [(r.name, r.price) for r in rows] == [("EDIFIER R1080 BT AUX / Bluetooth 5.0", 75)]
    gpu = {"raw_name": "MSI RTX 3060 б/у, без майнинга", "category": "gpu", "condition": "used", "price": 250}
    assert check_offer(gpu) == []
    assert [f["type"] for f in check_offer({**gpu, "raw_name": "MSI RTX 3060 после майнинга"})] == ["mining"]
