"""Builds the synthetic part of the Phase 7 benchmark (run once; needs F:/claudy/samples, which is outside git).

Every synthetic row starts from a clean row of the real Age Computers price list that the rule parser reads
completely (CPU, RAM, SSD, HDD, boards, video cards, monitors). The rule result of the clean row is the right answer;
the row is then made dirty (slang, typos, columns, merged cells, several tables, CSV cp1251, PDF, price formats).
Prices are shifted by a random 0-10% so the real price list of the shop is not published.

Usage (from backend/): .venv/Scripts/python.exe tests/benchmark/make_dataset.py
"""

import csv
import json
import random
import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND))

import openpyxl  # noqa: E402
from fpdf import FPDF  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402

from app.services.antifraud import check_offer  # noqa: E402
from app.services.catalog import make_match_key  # noqa: E402
from app.services.importer import file_to_markdown  # noqa: E402
from app.services.price_cleaner import clean_markdown  # noqa: E402
from app.services.rule_parser import parse_by_rules  # noqa: E402

AGE = Path(r"F:\claudy\samples\Price of Age Computers (9).xlsx")
OUT = Path(__file__).parent / "synthetic"
RATE = 11_850  # fixed UZS per USD for the whole benchmark
ROWS_PER_FILE = 150
FONT = r"C:\Windows\Fonts\arial.ttf"

TYPE_SLANG = {
    "cpu": ["проц", "процессор", "камень", "cpu"], "gpu": ["видюха", "видеокарта", "видяха", "vga"],
    "ram": ["озу", "оперативка", "оперативная память", "ram"], "ssd": ["ссд", "ssd", "твердотельник"],
    "hdd": ["винт", "хард", "жесткий диск", "hdd"], "motherboard": ["мать", "материнка", "мат. плата", "mb"],
    "monitor": ["моник", "монитор", "мон."],
}
BRAND_RU = {"intel": "интел", "amd": "амд", "kingston": "кингстон", "samsung": "самсунг", "gigabyte": "гигабайт",
            "asus": "асус", "msi": "мси", "adata": "адата", "lg": "лж", "toshiba": "тошиба", "palit": "палит"}
BRAND_TYPOS = {"gigabyte": "Gigabyet", "samsung": "Samsnug", "kingston": "Kingstom", "asus": "ASSUS", "intel": "Intell",
               "toshiba": "Toshiba", "msi": "MCI", "adata": "Adtata", "palit": "Pallit", "zotac": "Zotak",
               "asrock": "AsRok", "lexar": "Lexsar", "biostar": "Biostr"}
LOOKALIKE = {"K": "К", "E": "Е", "C": "С", "X": "Х", "T": "Т", "A": "А", "O": "О", "P": "Р", "H": "Н", "M": "М", "B": "В"}

# Rows with a known catch for Anti-Fraud and the spec ranges. flags: types expected; "typo" = a value out of range
TRAPS = [
    {"text": "Процессор Intel Core i9 (Xeon E5-2696 v3) 18 ядер 36 потоков LGA2011-3", "price": 45,
     "category": "cpu", "brand": "Intel", "model": "Xeon E5-2696 v3", "flags": ["server_hardware"], "trap": "xeon_as_i9"},
    {"text": "MB Huananzhi X99-TF + Xeon E5-2670v3 комплект", "price": 95, "category": "motherboard",
     "brand": "Huananzhi", "model": "X99-TF", "flags": ["server_hardware"], "trap": "china_x99"},
    {"text": "Samsung DDR4 16GB 2400MHz ECC REG", "price": 18, "category": "ram", "brand": "Samsung",
     "model": None, "flags": ["server_hardware"], "trap": "ecc_reg"},
    {"text": "Kingston DDR5 16GB 3200MHz (KVR32N22S8/16 DDR4)", "price": 41, "category": "ram", "brand": "Kingston",
     "model": "KVR32N22S8/16", "flags": ["spec_mismatch"], "trap": "ddr5_ddr4"},
    {"text": "Видеокарта MSI GTX 3060 Ventus 12GB", "price": 230, "category": "gpu", "brand": "MSI", "model": None,
     "flags": ["spec_mismatch"], "trap": "gtx_3060"},
    {"text": "Palit RTX 3070 GamingPro 8GB NO LHR", "price": 290, "category": "gpu", "brand": "Palit",
     "model": "RTX 3070 GamingPro", "flags": ["no_lhr"], "trap": "no_lhr"},
    {"text": "USB флешка Kingston DataTraveler 2TB", "price": 15, "category": "flash", "brand": "Kingston",
     "model": "DataTraveler", "flags": ["fake_capacity"], "trap": "flash_2tb"},
    {"text": "SSD M2 Samsung 4TB NVMe 990 Pro", "price": 40, "category": "ssd", "brand": "Samsung", "model": "990 Pro",
     "flags": ["fake_capacity"], "trap": "ssd_cheap_tb"},
    {"text": "GMOC DDR5 16GB 600MHz RGB", "price": 45, "category": "ram", "brand": "GMOC", "model": None,
     "flags": ["typo"], "trap": "typo_600mhz"},
    {"text": "Монитор LG 24\" 24GN60R IPS 1440Hz", "price": 160, "category": "monitor", "brand": "LG", "model": "24GN60R",
     "flags": ["typo"], "trap": "typo_1440hz"},
    {"text": "Блок питания Noname 600W", "price": 12, "category": "psu", "brand": None, "model": None,
     "flags": ["noname_psu"], "trap": "noname_psu"},
    {"text": "Видеокарта Sapphire RX 580 8GB после майнинга", "price": 60, "category": "gpu", "brand": "Sapphire",
     "model": "RX 580", "flags": ["mining"], "trap": "mining"},
    {"text": "Видеокарта P106-100 6GB без видеовыходов", "price": 35, "category": "gpu", "brand": None, "model": "P106-100",
     "flags": ["mining"], "trap": "p106"},
]
# Clean rows sold far below the market (the market = the clean price in two other stores)
BELOW_MARKET = {"s2_typos": 2, "s6_csv_cp1251": 1}
TRAP_FILES = {"s1_slang": TRAPS[0:5], "s2_typos": TRAPS[5:9], "s6_csv_cp1251": TRAPS[9:13]}


def load_clean() -> list[dict]:
    rows = clean_markdown(file_to_markdown(AGE.read_bytes(), AGE.name))
    result, seen = [], set()
    for r in rows:
        parsed = parse_by_rules(r.name)
        if not parsed or r.price is None or r.name in seen:
            continue
        seen.add(r.name)
        specs = {k: v for k, v in parsed["specs"].items() if not k.startswith("_")}
        item = {"source": r.name, "article": r.article, "clean_price": r.price, "category": parsed["category"],
                "brand": parsed["brand"], "model": parsed["model"], "name": parsed["name"], "specs": specs,
                "match_key": make_match_key(parsed["category"], parsed["name"], parsed["brand"], specs)}
        # A real catch already in the clean list (KingBank DDR5/DDR4, GTX 2060 Super) stays expected
        item["flags"] = sorted({f["type"] for f in check_offer(
            {**parsed, "raw_name": r.name, "article": r.article, "price": r.price, "currency": "USD"})})
        result.append(item)
    return result


def pick(clean: list[dict], rnd: random.Random, n: int) -> list[dict]:
    """n rows, every category in proportion (at least 5 rows each)."""
    by_cat: dict[str, list] = {}
    for item in clean:
        by_cat.setdefault(item["category"], []).append(item)
    share = {c: max(5, round(n * len(v) / len(clean))) for c, v in by_cat.items()}
    chosen = [x for c, v in by_cat.items() for x in rnd.sample(v, min(len(v), share[c]))]
    rnd.shuffle(chosen)
    return chosen[:n]


def label(item: dict, text: str, price: float, currency: str, rnd: random.Random, article: bool = False) -> dict:
    """The right answer of a dirty row. Expected flags: only what the dirty row itself shows (a slang row without
    the article "KJXS DDR4" has no visible DDR mismatch any more)."""
    keep = ("category", "brand", "model", "name", "match_key", "specs", "source", "article")
    written = {**item, "raw_name": text, "article": item["article"] if article else None, "price": price,
               "currency": currency}
    flags = sorted({f["type"] for f in check_offer(written, None, RATE)})
    return {"text": text, **{k: item[k] for k in keep}, "flags": flags, "price": price, "currency": currency,
            "condition": "new"}


def usd(item: dict, rnd: random.Random) -> int:
    return max(1, round(item["clean_price"] * rnd.uniform(0.9, 1.1)))


def uzs(price_usd: float) -> int:
    return round(price_usd * RATE, -3)


def slang(item: dict, rnd: random.Random) -> str:
    name = item["name"]
    name = re.sub(r"(\d+)GB\b", lambda m: rnd.choice(["{}гб", "{} гб", "{}gb", "{}г"]).format(m[1]), name)
    name = re.sub(r"(\d+)TB\b", lambda m: rnd.choice(["{}тб", "{} тб", "{}tb"]).format(m[1]), name)
    name = re.sub(r"(\d+)MHz\b", lambda m: rnd.choice(["{}мгц", "{}", "{}mhz"]).format(m[1]), name)
    name = re.sub(r"(\d+) rpm\b", lambda m: rnd.choice(["{}об", "{} rpm", "{}rpm"]).format(m[1]), name)
    name = re.sub(r'(\d+(?:\.\d)?)"', lambda m: rnd.choice(['{}"', "{} дюйм", "{}дюйма"]).format(m[1]), name)
    name = re.sub(r"\b(i[3579])-(\d)", r"\1 \2", name) if rnd.random() < 0.6 else name
    name = name.replace("GeForce ", "") if rnd.random() < 0.5 else name
    name = name.replace("Intel Core ", "Core ") if rnd.random() < 0.3 else name
    brand = (item["brand"] or "").lower()
    if brand in BRAND_RU and rnd.random() < 0.35:
        name = re.sub(re.escape(item["brand"]), BRAND_RU[brand], name, count=1, flags=re.IGNORECASE)
    if item["category"] == "cpu" and rnd.random() < 0.5:
        name = name.replace("OEM", rnd.choice(["оем", "tray", "oem"])).replace("BOX", rnd.choice(["бокс", "box"]))
    name = f"{rnd.choice(TYPE_SLANG[item['category']])} {name}"
    return name.lower() if rnd.random() < 0.5 else name


def typo(item: dict, rnd: random.Random) -> str:
    text = item["source"]
    if rnd.random() < 0.25:
        return text  # some rows stay clean in every file
    for _ in range(rnd.choice([1, 2])):
        kind = rnd.choice(["brand", "lookalike", "spaces", "units"])
        brand = (item["brand"] or "").lower()
        if kind == "brand" and brand:
            wrong = BRAND_TYPOS.get(brand)
            if not wrong and len(brand) > 3:
                i = rnd.randrange(1, len(brand) - 2)
                wrong = item["brand"][:i] + item["brand"][i + 1] + item["brand"][i] + item["brand"][i + 2:]
            if wrong:
                text = re.sub(re.escape(item["brand"]), wrong, text, count=1, flags=re.IGNORECASE)
        elif kind == "lookalike":
            codes = [m for m in re.finditer(r"\b(?=\w*\d)\w*[KECXTAOPHMB]\w*\b", text)]
            if codes:
                m = rnd.choice(codes)
                letters = [i for i, ch in enumerate(m[0]) if ch in LOOKALIKE]
                i = m.start() + rnd.choice(letters)
                text = text[:i] + LOOKALIKE[text[i]] + text[i + 1:]
        elif kind == "spaces":
            text = re.sub(r",\s", ",", text) if rnd.random() < 0.5 else re.sub(r"\s", "  ", text, count=2)
        else:
            text = text.replace("GHz", "Ггц").replace("MHz", "Mhz").replace("GB", "Gb").replace("Monitor", "Monitr")
    return text


def below_market(items: list[dict], rnd: random.Random, n: int, prices: dict) -> None:
    for item in rnd.sample(items, n):
        prices[item["text"]] = round(prices[item["text"]] * 0.35)
        item["price"] = prices[item["text"]]
        item["flags"] = sorted(set(item["flags"]) | {"price_below_market"})
        item["trap"] = "below_market"


def trap_labels(traps: list[dict]) -> list[dict]:
    return [{"text": t["text"], "category": t["category"], "brand": t["brand"], "model": t["model"], "name": None,
             "match_key": None, "specs": {}, "flags": t["flags"], "price": t["price"], "currency": "USD",
             "condition": "used" if "майнинг" in t["text"] else "new", "trap": t["trap"]} for t in traps]


def save_xlsx(wb, name: str) -> None:
    wb.save(OUT / f"{name}.xlsx")


def build(clean: list[dict]) -> None:
    OUT.mkdir(exist_ok=True)
    meta = {}

    # s1: slang and abbreviations, prices "110$" / "110 у.е." / plain
    rnd = random.Random(1)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels = []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Наименование", "Цена"])
    for item in items:
        p = usd(item, rnd)
        text = slang(item, rnd)
        ws.append([text, rnd.choice([f"{p}$", f"{p} у.е.", p])])
        labels.append(label(item, text, p, "USD", rnd))
    for t in TRAP_FILES["s1_slang"]:
        ws.append([t["text"], t["price"]])
    labels += trap_labels(TRAP_FILES["s1_slang"])
    save_xlsx(wb, "s1_slang")
    meta["s1_slang"] = {"file": "s1_slang.xlsx", "currency": "USD", "what": "сленг и сокращения, цены «110$», «110 у.е.»"}
    write_labels("s1_slang", labels)

    # s2: typos, Cyrillic look-alikes in codes, broken spaces; two rows far below the market
    rnd = random.Random(2)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels, prices = [], {}
    for item in items:
        text = typo(item, rnd)
        prices[text] = usd(item, rnd)
        labels.append(label(item, text, prices[text], "USD", rnd, article=True))
    below_market(labels[:60], rnd, BELOW_MARKET["s2_typos"], prices)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Наименование", "Артикул", "Цена"])
    for lab, item in zip(labels, items):
        ws.append([lab["text"], item["article"], prices[lab["text"]]])
    for t in TRAP_FILES["s2_typos"]:
        ws.append([t["text"], None, t["price"]])
    labels += trap_labels(TRAP_FILES["s2_typos"])
    save_xlsx(wb, "s2_typos")
    meta["s2_typos"] = {"file": "s2_typos.xlsx", "currency": "USD", "what": "опечатки в брендах, кириллица в кодах, пробелы"}
    write_labels("s2_typos", labels)

    # s3: mixed-up and nameless columns; the second sheet has no header at all
    rnd = random.Random(3)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels = []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Цена, $", "Артикул", None, "Наименование товара", "Наличие"])
    ws2 = wb.create_sheet("Лист2")
    for i, item in enumerate(items):
        p = usd(item, rnd)
        if i < len(items) // 2:
            ws.append([p, item["article"], rnd.choice(["", "new", "хит"]), item["source"], rnd.choice(["есть", "в наличии", "+"])])
        else:
            ws2.append([i + 1, "есть", item["source"], p])
        labels.append(label(item, item["source"], p, "USD", rnd, article=i < len(items) // 2))
    save_xlsx(wb, "s3_columns")
    meta["s3_columns"] = {"file": "s3_columns.xlsx", "currency": "USD", "what": "столбцы в другом порядке, безымянные, лист без шапки"}
    write_labels("s3_columns", labels)

    # s4: merged cells: a category title over the table width, a brand cell merged down over its rows; UZS
    rnd = random.Random(4)
    items = sorted(pick(clean, rnd, ROWS_PER_FILE), key=lambda x: (x["category"], x["brand"] or ""))
    labels = []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Бренд", "Модель", "Цена, сум"])
    titles = {"cpu": "ПРОЦЕССОРЫ", "gpu": "ВИДЕОКАРТЫ", "ram": "ОПЕРАТИВНАЯ ПАМЯТЬ", "ssd": "SSD НАКОПИТЕЛИ",
              "hdd": "ЖЁСТКИЕ ДИСКИ", "motherboard": "МАТЕРИНСКИЕ ПЛАТЫ", "monitor": "МОНИТОРЫ"}
    category, start, brand = None, None, None
    for item in items:
        if item["category"] != category:
            if start and ws.max_row > start:
                ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
            category, brand, start = item["category"], None, None
            ws.append([titles[category]])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        model = re.sub(rf"^{re.escape(item['brand'] or '')}\W*", "", item["source"], flags=re.IGNORECASE)
        p = usd(item, rnd)
        if item["brand"] != brand:
            if start and ws.max_row > start:
                ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
            brand, start = item["brand"], ws.max_row + 1
        ws.append([item["brand"] if ws.max_row + 1 == start else None, model, uzs(p)])
        ws.cell(ws.max_row, 1).alignment = Alignment(vertical="center")
        labels.append(label(item, model, uzs(p), "UZS", rnd))
    if start and ws.max_row > start:
        ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
    save_xlsx(wb, "s4_merged")
    meta["s4_merged"] = {"file": "s4_merged.xlsx", "currency": "UZS", "what": "объединённые ячейки: заголовок раздела, бренд на несколько строк"}
    write_labels("s4_merged", labels)

    # s5: several tables on one sheet: two "name | price" side by side, then one more table below; UZS strings
    rnd = random.Random(5)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels = []
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Оплата только в национальной валюте!!!"])
    ws.append([])
    left, right, below = items[:55], items[55:110], items[110:]
    ws.append(["Наименование", "Цена", None, "Наименование", "Цена"])
    for i in range(max(len(left), len(right))):
        cells = [None] * 5
        for col, part in ((0, left), (3, right)):
            if i < len(part):
                p = uzs(usd(part[i], rnd))
                cells[col], cells[col + 1] = part[i]["source"], rnd.choice([f"{p:,}".replace(",", " "), f"{p:,}".replace(",", ".") + " сум"])
                labels.append(label(part[i], part[i]["source"], p, "UZS", rnd))
        ws.append(cells)
    ws.append([])
    ws.append(["Новое поступление"])
    ws.append(["№", "Наименование", "Артикул", "Цена"])
    for i, item in enumerate(below):
        p = uzs(usd(item, rnd))
        ws.append([i + 1, item["source"], item["article"], f"{p:,}".replace(",", " ")])
        labels.append(label(item, item["source"], p, "UZS", rnd, article=True))
    save_xlsx(wb, "s5_multitable")
    meta["s5_multitable"] = {"file": "s5_multitable.xlsx", "currency": "UZS",
                             "what": "несколько таблиц на листе (две рядом, одна ниже), цены «1 250 000», «1.250.000 сум»"}
    write_labels("s5_multitable", labels)

    # s6: CSV with ";" in cp1251, mixed price formats and currencies in one column
    rnd = random.Random(6)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels, prices = [], {}
    for item in items:
        prices[item["source"]] = usd(item, rnd)
        labels.append(label(item, item["source"], prices[item["source"]], "USD", rnd))
    below_market(labels[:60], rnd, BELOW_MARKET["s6_csv_cp1251"], prices)
    with open(OUT / "s6_csv_cp1251.csv", "w", encoding="cp1251", errors="replace", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Наименование", "Цена"])
        for lab in labels:
            p = prices[lab["text"]]
            fmt = rnd.choice(["sum_spaces", "sum_dots", "dollar", "ye", "comma"])
            if fmt in ("sum_spaces", "sum_dots"):
                lab["price"], lab["currency"] = uzs(p), "UZS"
                cell = f"{uzs(p):,}".replace(",", " " if fmt == "sum_spaces" else ".") + " сум"
            else:
                cell = {"dollar": f"{p}$", "ye": f"{p} у.е.", "comma": f"{p},00"}[fmt]
            w.writerow([lab["text"], cell])
        for t in TRAP_FILES["s6_csv_cp1251"]:
            w.writerow([t["text"], f"{t['price']}$"])
    labels += trap_labels(TRAP_FILES["s6_csv_cp1251"])
    meta["s6_csv_cp1251"] = {"file": "s6_csv_cp1251.csv", "currency": "USD",
                             "what": "CSV с «;» в cp1251, цены «1 250 000 сум», «1.250.000 сум», «110$», «110 у.е.», «110,00»"}
    write_labels("s6_csv_cp1251", labels)

    # s7: PDF table, prices "110$" and "1 250 000 сум"
    rnd = random.Random(7)
    items = pick(clean, rnd, ROWS_PER_FILE)
    labels = []
    pdf = FPDF(orientation="L", format="A4")
    pdf.add_font("arial", fname=FONT)
    pdf.add_font("arial", style="B", fname=FONT.replace("arial.ttf", "arialbd.ttf"))
    pdf.set_font("arial", size=8)
    pdf.add_page()
    pdf.cell(0, 8, "Прайс-лист компьютерной техники", new_x="LMARGIN", new_y="NEXT")
    with pdf.table(col_widths=(200, 40), text_align=("LEFT", "RIGHT")) as table:
        table.row(["Наименование", "Цена"])
        for item in items:
            p = usd(item, rnd)
            if rnd.random() < 0.3:
                cell, price, currency = f"{uzs(p):,}".replace(",", " ") + " сум", uzs(p), "UZS"
            else:
                cell, price, currency = f"{p}$", p, "USD"
            table.row([item["source"], cell])
            labels.append(label(item, item["source"], price, currency, rnd))
    pdf.output(str(OUT / "s7_pdf.pdf"))
    meta["s7_pdf"] = {"file": "s7_pdf.pdf", "currency": "USD", "what": "PDF-таблица, цены «110$» и «1 250 000 сум»"}
    write_labels("s7_pdf", labels)

    (OUT / "files.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    # Market for "price below market": the clean price of every reference card in two other stores
    rnd, market = random.Random(8), {}
    for item in clean:
        base = item["clean_price"] * rnd.uniform(0.95, 1.05)
        market[item["match_key"]] = [round(base, 2), round(base * 1.04, 2)]
    (OUT / "market.json").write_text(json.dumps(market, ensure_ascii=False), encoding="utf-8")


def write_labels(name: str, labels: list[dict]) -> None:
    for i, lab in enumerate(labels, 1):
        lab["id"] = f"{name[:2]}-{i:03d}"
    (OUT / f"{name}.json").write_text(json.dumps(labels, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    build(load_clean())
